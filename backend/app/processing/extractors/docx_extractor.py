"""DOCX extractor — python-docx, document order preserved.

Walks the document body in order, grouping consecutive paragraphs into a
section and emitting each table with headers/rows preserved as data.
No semantic interpretation, no OCR.
"""

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.constants import PageContentType, TextExtractionStatus
from app.processing.base import (
    DocumentExtractor,
    ExtractedSection,
    ExtractionResult,
    package_version,
)
from app.processing.normalization import normalize_text

EXTRACTOR_NAME = "python-docx"


def _table_to_structured(table: Table, table_index: int) -> dict:
    """Convert a DOCX table to {table_index, headers, rows} without type guessing."""
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    headers = rows[0] if rows else []
    return {
        "table_index": table_index,
        "headers": headers,
        "rows": rows[1:] if rows else [],
    }


class DOCXExtractor(DocumentExtractor):
    name = EXTRACTOR_NAME
    version = package_version("python-docx")

    def extract(self, file) -> ExtractionResult:
        result = ExtractionResult(
            extractor_name=self.name, extractor_version=self.version
        )

        document = DocxDocument(file)

        # Iterate body elements in true document order (python-docx's
        # document.paragraphs / document.tables lose interleaving).
        body = document.element.body
        paragraph_map = {p._p: p for p in document.paragraphs}
        table_map = {t._tbl: t for t in document.tables}

        paragraph_lines: list[str] = []
        table_index = 0
        section_index = 0

        def flush_paragraphs() -> None:
            nonlocal section_index
            if not paragraph_lines:
                return
            section_index += 1
            text = normalize_text("\n".join(paragraph_lines))
            result.sections.append(
                ExtractedSection(
                    index=section_index,
                    content_type=PageContentType.PAGE,
                    section_reference=f"section {section_index}",
                    text=text or None,
                    extraction_status=(
                        TextExtractionStatus.EXTRACTED if text.strip() else TextExtractionStatus.NO_TEXT
                    ),
                    structured_metadata=None,
                )
            )
            paragraph_lines.clear()

        for child in body.iterchildren():
            if child in paragraph_map:
                paragraph = paragraph_map[child]
                style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
                prefix = ""
                if style.startswith("heading"):
                    prefix = f"[{paragraph.style.name}] "
                if paragraph.text.strip():
                    paragraph_lines.append(f"{prefix}{paragraph.text}")
                else:
                    paragraph_lines.append("")
            elif child in table_map:
                flush_paragraphs()  # keep reading order: paragraphs, then this table
                table_index += 1
                structured = _table_to_structured(table_map[child], table_index)
                text = normalize_text(
                    "\n".join("\t".join(str(cell) for cell in row) for row in [structured["headers"], *structured["rows"]])
                )
                section_index += 1
                result.sections.append(
                    ExtractedSection(
                        index=section_index,
                        content_type=PageContentType.PAGE,
                        section_reference=f"section {section_index}, table {table_index}",
                        text=text or None,
                        extraction_status=(
                            TextExtractionStatus.EXTRACTED if text.strip() else TextExtractionStatus.NO_TEXT
                        ),
                        structured_metadata={"tables": [structured]},
                    )
                )

        flush_paragraphs()

        if not result.sections:
            raise ValueError("DOCX contains no readable content.")
        return result
