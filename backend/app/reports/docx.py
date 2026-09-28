"""DOCX report writer (Step 11) — python-docx rendering of deterministic sections.

Turns the built sections (app/reports/sections.py) into a clean .docx
document: title block, reporting period, generated timestamp, headings and
tables. Every rendered value originates from the collected structured
records — the writer adds formatting only, never data.

DOCX is the first supported output because python-docx is already a project
dependency (Step 4 extraction). PDF is deliberately not offered — the
project has no verified PDF generation mechanism.

The artifact is produced entirely in memory (BytesIO); nothing is written to
the source-of-truth storage and no database row is created (see limitation
notes in the API layer).
"""

import io
import zipfile
from datetime import timezone
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Pt

from app.reports.engine import REPORT_DATA_MISSING, ReportData
from app.reports.sections import PROTOTYPE_NOTICE, build_sections

TABLE_STYLE = "Table Grid"


def _add_table(document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = TABLE_STYLE
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


def _add_section(document, section: dict[str, Any]) -> None:
    document.add_heading(section.get("title", "Section"), level=1)

    for paragraph in section.get("paragraphs", []):
        document.add_paragraph(paragraph)
    for note in section.get("notes", []):
        paragraph = document.add_paragraph(note)
        for run in paragraph.runs:
            run.font.italic = True

    table = section.get("table")
    if table and table.get("rows"):
        _add_table(document, table["headers"], table["rows"])
    if section.get("headers") and section.get("rows"):
        _add_table(document, section["headers"], section["rows"])
    for grouped in section.get("tables", []):
        document.add_heading(grouped.get("caption", ""), level=2)
        _add_table(document, grouped["headers"], grouped["rows"])

    for item in section.get("items", []):
        headline = document.add_paragraph()
        run = headline.add_run(item.get("headline", ""))
        run.font.bold = True
        status = item.get("status")
        if status:
            status_run = headline.add_run(f"  [{status}]")
            status_run.font.bold = True
        for line in item.get("lines", []):
            document.add_paragraph(line, style="List Bullet")

    if section.get("review_required"):
        banner = document.add_paragraph(
            "REVIEW REQUIRED: this section contains conflicting values that "
            "were deliberately NOT resolved — a human must decide."
        )
        for run in banner.runs:
            run.font.bold = True

    if section.get("empty_note"):
        document.add_paragraph(section["empty_note"])


def render_report_docx(data: ReportData) -> bytes:
    """Render the collected report data into DOCX bytes (deterministic
    content; only the timestamp block varies between runs)."""
    spec = data.specification
    document = DocxDocument()

    document.add_heading(spec.title, level=0)

    meta = document.add_paragraph()
    generated = spec.generated_at.astimezone(timezone.utc)
    meta_lines = [
        f"Report type: {spec.report_type}",
        f"Reporting period: {spec.reporting_period or 'all available periods'}",
        f"Generated at: {generated.strftime('%Y-%m-%d %H:%M:%S UTC')}",
        f"Report id: {spec.report_id()}",
    ]
    if spec.requester:
        meta_lines.append(f"Requested by: {spec.requester}")
    if spec.document_ids:
        meta_lines.append(f"Selected documents: {', '.join(map(str, spec.document_ids))}")
    for line in meta_lines:
        meta.add_run(line + "\n")

    notice = document.add_paragraph(PROTOTYPE_NOTICE)
    for run in notice.runs:
        run.font.italic = True

    document.add_paragraph(
        "All values below originate from validated structured records stored "
        "by this platform. Values marked "
        f"'{REPORT_DATA_MISSING}' were absent from the source data and were "
        "NOT estimated. Conflicting values are reported with their sources; "
        "none was resolved automatically."
    )

    for section in build_sections(data):
        _add_section(document, section)

    buffer = io.BytesIO()
    document.save(buffer)
    return _canonical_zip(buffer)


# Fixed container timestamp: identical specs + identical data produce
# byte-identical artifacts regardless of when they were generated (the DOCX
# zip otherwise embeds wall-clock file times with 2-second granularity).
_FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def _canonical_zip(buffer: io.BytesIO) -> bytes:
    """Re-zip the DOCX with fixed entry timestamps (deterministic bytes)."""
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as source:
        entries = [
            (item.filename, source.read(item.filename), item.compress_type)
            for item in source.infolist()
        ]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for name, data, compress_type in entries:
            info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_TIMESTAMP)
            info.compress_type = compress_type
            target.writestr(info, data)
    return output.getvalue()


def report_filename(spec) -> str:
    """Deterministic artifact filename from fingerprint + timestamp."""
    stamp = spec.generated_at.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return f"{spec.report_id()}-{stamp}.docx"
