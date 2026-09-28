"""Report generation API integration tests — REQUIRE PostgreSQL (conftest
skip contract: these SKIP, never fake-pass, when the database is absent).

Covers the HTTP contract over real data: JSON metadata mode, DOCX attachment
mode, audit logging, and honest empty-store behavior.
"""

import io
from datetime import datetime, timezone

import pytest
from docx import Document as DocxDocument
from sqlalchemy.orm import Session

pytestmark = pytest.mark.db


def _make_processed_document(migrated_engine, filename: str, rows: list[dict]) -> int:
    """Insert a processed document with records (source-of-truth rows)."""
    from app.models import Document, ExtractedRecord

    with Session(migrated_engine) as session:
        document = Document(
            filename=filename,
            document_type="xlsx",
            storage_reference=f"test/{filename}",
            status="processed",
            processed_at=datetime.now(timezone.utc),
        )
        session.add(document)
        session.flush()
        for row in rows:
            session.add(ExtractedRecord(document_id=document.id, **row))
        session.commit()
        return document.id


def _generate(client, json_body=None, **overrides):
    body = {
        "title": "Integration Production Report",
        "report_type": "production",
        "reporting_period": "2025-06-30",
    }
    body.update(overrides)
    if json_body is not None:
        body.update(json_body)
    return client.post("/api/reports/generate", json=body)


def test_generate_json_mode_returns_metadata(client, migrated_engine):
    _make_processed_document(
        migrated_engine,
        "report.xlsx",
        [
            dict(record_type="coal_production", entity_name="DEMO_MINE_A",
                 metric_name="demo_coal_production", metric_value=1200,
                 value_raw="1,200", normalized_value="1200", unit="tonnes",
                 reporting_period="2025-06-30", extraction_method="spreadsheet",
                 validation_status="pass", source_reference="Sheet1!B3"),
        ],
    )
    response = _generate(client, {"output_format": "json"})
    assert response.status_code == 200
    payload = response.json()
    report = payload["report"] if "report" in payload else payload
    assert report["status"] == "ok"
    assert report["evidence_count"] == 1
    assert report["conflict_count"] == 0
    assert report["artifact"]["format"] == "json"
    assert report["artifact"]["filename"].endswith(".docx")
    assert report["specification"]["title"] == "Integration Production Report"


def test_generate_docx_attachment_with_metadata_headers(client, migrated_engine):
    _make_processed_document(
        migrated_engine,
        "report2.xlsx",
        [
            dict(record_type="coal_production", entity_name="DEMO_MINE_A",
                 metric_name="demo_coal_production", metric_value=1200,
                 value_raw="1200", normalized_value="1200", unit="tonnes",
                 reporting_period="2025-06-30", validation_status="pass"),
        ],
    )
    response = _generate(client)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.headers["x-report-records"] == "1"
    assert response.headers["x-report-conflicts"] == "0"
    document = DocxDocument(io.BytesIO(response.content))
    assert any(p.text == "Integration Production Report" for p in document.paragraphs)


def test_generate_records_audit_row_with_metadata_only(client, migrated_engine):
    from sqlalchemy import select

    from app.models import AuditLog

    _make_processed_document(
        migrated_engine,
        "report3.xlsx",
        [
            dict(record_type="coal_production", entity_name="DEMO_MINE_A",
                 metric_name="demo_coal_production", metric_value=1200,
                 value_raw="1200", normalized_value="1200", unit="tonnes",
                 reporting_period="2025-06-30", validation_status="pass"),
        ],
    )
    response = _generate(client, {"output_format": "json"})
    assert response.status_code == 200
    with Session(migrated_engine) as session:
        rows = list(
            session.scalars(select(AuditLog).where(AuditLog.action == "report.generate"))
        )
        assert len(rows) == 1
        details = rows[0].details or {}
        assert details["status"] == "ok"
        assert details["record_count"] == 1
        flat = str(details)
        assert "DEMO_MINE_A" not in flat and "1200" not in flat


def test_generate_empty_store_is_honest_not_error(client):
    response = _generate(client, {"output_format": "json"})
    assert response.status_code == 200
    report = response.json()["report"] if "report" in response.json() else response.json()
    assert report["status"] == "ok"
    assert report["evidence_count"] == 0
    assert report["summary"]["record_count"] == 0


def test_generate_rejects_invalid_spec_with_422(client):
    response = _generate(client, {"output_format": "pdf"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "unsupported_report_format"
    response = _generate(client, {"sections": ["made_up_section"]})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_sections"


def test_generate_reports_conflicts_without_resolution(client, migrated_engine):
    doc_a = _make_processed_document(
        migrated_engine, "a.xlsx",
        [dict(record_type="coal_production", entity_name="MINE_X",
              metric_name="production", metric_value=1200, value_raw="1200",
              normalized_value="1200", unit="tonnes",
              reporting_period="2025-06-30", validation_status="pass")],
    )
    _make_processed_document(
        migrated_engine, "b.xlsx",
        [dict(record_type="coal_production", entity_name="MINE_X",
              metric_name="production", metric_value=1350, value_raw="1350",
              normalized_value="1350", unit="tonnes",
              reporting_period="2025-06-30", validation_status="pass")],
    )
    response = _generate(client, {"output_format": "json", "report_type": "production"})
    assert response.status_code == 200
    report = response.json()["report"] if "report" in response.json() else response.json()
    assert report["conflict_count"] == 1
    values = [v["value_raw"] for v in report["conflicts"][0]["values"]]
    assert values == ["1200", "1350"]  # both preserved, no winner
    assert report["conflicts"][0]["status"] == "REVIEW REQUIRED"
    assert doc_a > 0
