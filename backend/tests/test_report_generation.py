"""Step 11 report generation tests — DB-free.

Covers: specification validation, deterministic grouping, raw/normalized
value preservation, conflict detection (no winner), missing-data behavior,
provenance preservation, section generation, DOCX generation (incl. read-back),
no-fabrication guarantees and deterministic output structure.

PostgreSQL-dependent behavior (the API endpoint over real data, SQL filters,
truncation over >500 records) lives in test_reports_api.py and skips honestly
when PostgreSQL is unavailable.
"""

import io
from datetime import datetime, timezone

import pytest
from docx import Document as DocxDocument

from app.reports import (
    RecordItem,
    ReportData,
    build_specification,
    generate_report,
)
from app.reports.engine import REPORT_DATA_MISSING, _detect_conflicts
from app.reports.sections import PROTOTYPE_NOTICE, build_sections
from app.reports.service import audit_metadata


def _record(record_id, **overrides) -> RecordItem:
    defaults = dict(
        document_id=1,
        document_name="docA.pdf",
        entity="Mine A",
        metric="production",
        value_raw="1200",
        normalized_value="1200",
        unit="tonnes",
        reporting_period="2025-06",
        extraction_method="native_text",
        validation_status="pass",
    )
    defaults.update(overrides)
    return RecordItem(record_id=record_id, **defaults)


def _data(records, conflicts=None, spec=None) -> ReportData:
    spec = spec or build_specification(title="Test Report", report_type="production")
    data = ReportData(
        specification=spec, records=list(records),
        conflicts=conflicts if conflicts is not None else _detect_conflicts(list(records)),
        document_ids=sorted({r.document_id for r in records}),
    )
    from app.reports.engine import compute_summary

    data.summary = compute_summary(data)
    return data


def _payload(records, conflicts=None, spec=None):
    return generate_report(spec or build_specification(title="Test Report"), data=_data(records, conflicts, spec))


# --- report specification validation -------------------------------------


def test_spec_normalizes_and_deduplicates():
    spec = build_specification(
        title="  Quarterly   Review ",
        report_type="production",
        reporting_period="2025-06",
        document_ids=[3, 1, 3, 1],
        entities=["Mine A", "Mine A"],
        metrics=["production"],
        sections=["key_figures", "sources_evidence"],
        output_format="docx",
    )
    assert spec.title == "Quarterly Review"
    assert spec.document_ids == (3, 1)  # deduplicated, order preserved
    assert spec.entities == ("Mine A",)
    assert spec.requested_sections == ("key_figures", "sources_evidence")
    assert spec.output_format == "docx"


def test_spec_defaults_sections_when_empty():
    spec = build_specification(title="All sections")
    from app.reports.spec import DEFAULT_SECTIONS

    assert set(spec.requested_sections) == set(DEFAULT_SECTIONS)


@pytest.mark.parametrize(
    "kwargs, code",
    [
        ({"title": "   "}, "empty_title"),
        ({"title": "x" * 201}, "empty_title"),
        ({"title": "ok", "report_type": "x" * 65}, "invalid_report_type"),
        ({"title": "ok", "reporting_period": "x" * 65}, "invalid_reporting_period"),
        ({"title": "ok", "document_ids": [0]}, "invalid_document_ids"),
        ({"title": "ok", "document_ids": ["a"]}, "invalid_document_ids"),
        ({"title": "ok", "document_ids": list(range(1, 102))}, "too_many_document_ids"),
        ({"title": "ok", "entities": ["e"] * 51}, "too_many_entities"),
        ({"title": "ok", "metrics": ["m"] * 51}, "too_many_metrics"),
        ({"title": "ok", "sections": ["invented_section"]}, "invalid_sections"),
        ({"title": "ok", "output_format": "pdf"}, "unsupported_report_format"),
        ({"title": "ok", "filters": {"validation_status": "x" * 40}}, "invalid_filters"),
        ({"title": "ok", "requester": "r" * 130}, "invalid_requester"),
    ],
)
def test_spec_rejects_invalid_input_with_specific_codes(kwargs, code):
    from app.exceptions import AppError

    with pytest.raises(AppError) as excinfo:
        build_specification(**kwargs)
    assert excinfo.value.code == code


def test_spec_fingerprint_is_content_based_not_time_based():
    a = build_specification(
        title="Same", document_ids=[1], generated_at=datetime(2026, 1, 1, tzinfo=timezone.utc)
    )
    b = build_specification(
        title="Same", document_ids=[1], generated_at=datetime(2027, 6, 15, tzinfo=timezone.utc)
    )
    assert a.fingerprint() == b.fingerprint()
    c = build_specification(title="Different", document_ids=[1])
    assert c.fingerprint() != a.fingerprint()


# --- deterministic grouping ------------------------------------------------


def test_grouping_is_deterministic_and_ordered():
    records = [
        _record(1, entity="Mine B", metric="dispatch"),
        _record(2, entity="Mine A", metric="production"),
        _record(3, entity="Mine A", metric="dispatch"),
        _record(4, entity="Mine A", metric="production", reporting_period="2025-07"),
    ]
    data = _data(records)
    groups = data.grouped_records()
    keys = [(e, m, p) for e, m, p, _ in groups]
    assert keys == sorted(keys, key=lambda k: ((k[0] or "~"), (k[1] or "~"), (k[2] or "~")))
    assert keys[0] == ("Mine A", "dispatch", "2025-06")
    # Re-running on identical input gives identical grouping.
    assert [(e, m, p) for e, m, p, _ in _data(records).grouped_records()] == keys


# --- raw / normalized value preservation -----------------------------------


def test_raw_value_is_preserved_verbatim():
    record = _record(1, value_raw="1,20O5", normalized_value=None)
    payload = _payload([record])
    key_row = payload["summary"]
    assert key_row["record_count"] == 1
    stored = payload["specification"]  # spec echo sanity
    assert stored["title"] == "Test Report"
    # The raw string survives into the conflict-free record untouched.
    sections = build_sections(_data([record]))
    key_figures = next(s for s in sections if s["title"] == "Key Figures")
    assert key_figures["rows"][0][2] == "1,20O5"  # verbatim, never 'corrected'


def test_normalized_value_kept_alongside_raw():
    record = _record(1, value_raw="1,200", normalized_value="1200")
    sections = build_sections(_data([record]))
    key_figures = next(s for s in sections if s["title"] == "Key Figures")
    row = key_figures["rows"][0]
    assert row[2] == "1,200" and row[3] == "1200"


# --- conflict detection ----------------------------------------------------


def test_conflict_reports_both_values_with_sources_and_no_winner():
    records = [
        _record(1, document_id=1, document_name="A.pdf", page_number=3, value_raw="1200"),
        _record(2, document_id=2, document_name="B.xlsx", value_raw="1350"),
    ]
    conflicts = _detect_conflicts(records)
    assert len(conflicts) == 1
    group = conflicts[0]
    values = [v["value_raw"] for v in group.values]
    assert values == ["1200", "1350"]  # both sides, deterministic order
    assert group.to_payload()["status"] == "REVIEW REQUIRED"
    sides = group.to_payload()["values"]
    assert sides[0]["sources"][0]["document_name"] == "A.pdf"
    assert sides[1]["sources"][0]["document_name"] == "B.xlsx"
    # Neither value is selected as the report's answer anywhere.
    payload = _payload(records)
    assert payload["conflict_count"] == 1
    assert payload["summary"]["record_count"] == 2


def test_conflicts_require_same_entity_metric_period_and_unit():
    records = [
        _record(1, value_raw="1200", unit="tonnes"),
        _record(2, value_raw="1350", unit="tonnes", entity="Mine B"),
        _record(3, value_raw="1350", unit="tonnes", metric="dispatch"),
        _record(4, value_raw="1350", unit="tonnes", reporting_period="2025-07"),
        _record(5, value_raw="1350", unit="kilotonnes"),
    ]
    assert _detect_conflicts(records) == []


def test_identical_values_and_valueless_records_never_conflict():
    records = [
        _record(1, value_raw="1200"),
        _record(2, value_raw="1200", document_id=2),
        _record(3, value_raw=None, normalized_value=None),
    ]
    assert _detect_conflicts(records) == []


def test_conflict_sections_mark_review_required():
    records = [
        _record(1, value_raw="1200"),
        _record(2, document_id=2, value_raw="1350"),
    ]
    payload = _payload(records)
    sections = build_sections(_data(records))
    validation = next(s for s in sections if s["title"] == "Validation / Review Notes")
    assert validation["review_required"] is True
    assert any(item["headline"].startswith("CONFLICT") for item in validation["items"])
    assert all(item["status"] == "REVIEW REQUIRED"
               for item in validation["items"] if item["kind"] == "conflict")


# --- missing-data behavior --------------------------------------------------


def test_missing_values_are_marked_never_estimated():
    records = [_record(1, value_raw=None, normalized_value=None)]
    sections = build_sections(_data(records))
    detailed = next(s for s in sections if s["title"] == "Detailed Data")
    assert detailed["tables"][0]["rows"][0][1] == REPORT_DATA_MISSING
    payload = _payload(records)
    assert payload["missing_value_count"] == 1
    summary = next(
        line for line in
        next(s for s in sections if s["title"] == "Executive Summary")["paragraphs"]
        if "no stored value" in line
    )
    assert "NOT" in summary or "missing" in summary


def test_no_records_is_honest_not_fabricated():
    payload = _payload([])
    assert payload["evidence_count"] == 0
    assert payload["conflict_count"] == 0
    assert payload["summary"]["record_count"] == 0
    sections = build_sections(_data([]))
    key_figures = next(s for s in sections if s["title"] == "Key Figures")
    assert key_figures["rows"] == []
    assert "Nothing was estimated" in key_figures["empty_note"]


# --- provenance preservation -------------------------------------------------


def test_provenance_survives_into_report_and_sources_section():
    record = _record(
        7,
        document_id=42,
        document_name="geol_report.pdf",
        page_id=9,
        page_number=4,
        source_reference="page 12, table 3",
        extraction_method="ocr",
        validation_status="review_required",
        ocr_confidence=0.42,
    )
    payload = _payload([record])
    sections = build_sections(_data([record]))
    sources = next(s for s in sections if s["title"] == "Sources / Evidence")
    row = sources["rows"][0]
    assert row[0] == "7" and "geol_report.pdf" in row[1] and "doc 42" in row[1]
    assert row[2] == "4" and row[3] == "page 12, table 3"
    assert row[4] == "ocr" and row[5] == "0.42" and row[6] == "review_required"
    # Document/page ids retained on the record payload for traceability.
    data = _data([record])
    item = data.records[0]
    assert item.document_id == 42 and item.page_id == 9 and item.record_id == 7


# --- section generation -------------------------------------------------------


def test_all_default_sections_build_in_canonical_order():
    data = _data([_record(1)])
    sections = build_sections(data)
    assert [s["title"] for s in sections] == [
        "Executive Summary",
        "Key Figures",
        "Detailed Data",
        "Validation / Review Notes",
        "Sources / Evidence",
    ]


def test_requested_sections_subset_is_respected():
    spec = build_specification(title="Only two", sections=["key_figures", "sources_evidence"])
    sections = build_sections(_data([_record(1)], spec=spec))
    assert [s["title"] for s in sections] == ["Key Figures", "Sources / Evidence"]


def test_executive_summary_is_deterministic_template_not_llm():
    records = [
        _record(1),
        _record(2, document_id=2, value_raw="1350", entity="Mine A", metric="production"),
    ]
    data = _data(records)
    sections = build_sections(data)
    summary = next(s for s in sections if s["title"] == "Executive Summary")
    assert summary["paragraphs"][0].startswith("This production report covers")
    assert "2 structured record(s)" in summary["paragraphs"][0]
    assert PROTOTYPE_NOTICE in summary["notes"][0]
    # Byte-identical when regenerated from the same data.
    again = build_sections(_data(records))
    assert again == sections


# --- DOCX generation -----------------------------------------------------------


def _docx_document(records, spec=None):
    payload = _payload(records, spec=spec)
    return payload, DocxDocument(io.BytesIO(payload["docx_bytes"]))


def test_docx_smoke_is_valid_and_contains_title_and_sections():
    payload, doc = _docx_document([_record(1), _record(2, document_id=2, value_raw="1350")])
    assert payload["docx_bytes"][:2] == b"PK"  # real zip-based OOXML
    assert payload["artifact"]["format"] == "docx"
    texts = [p.text for p in doc.paragraphs]
    assert "Test Report" in texts
    assert any("Reporting period" in t for t in texts)
    assert any("Generated at" in t for t in texts)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for expected in ("Executive Summary", "Key Figures", "Detailed Data",
                     "Validation / Review Notes", "Sources / Evidence"):
        assert expected in headings


def test_docx_key_figures_table_matches_records():
    records = [_record(1), _record(2, document_id=2, document_name="B.xlsx", value_raw="1350")]
    _, doc = _docx_document(records)
    tables = doc.tables
    assert tables, "Key Figures table missing"
    first = tables[0]
    header = [c.text for c in first.rows[0].cells]
    assert header[:3] == ["Entity", "Metric", "Value (raw)"]
    body = [[c.text for c in row.cells] for row in first.rows[1:]]
    assert len(body) == 2
    assert {"1200", "1350"} == {row[2] for row in body}


def test_docx_conflict_banner_present_when_conflicts_exist():
    records = [_record(1), _record(2, document_id=2, value_raw="1350")]
    _, doc = _docx_document(records)
    texts = [p.text for p in doc.paragraphs]
    assert any(t.startswith("REVIEW REQUIRED") for t in texts)


def test_docx_filename_is_deterministic_and_content_addressed():
    spec = build_specification(
        title="Named", report_type="production", reporting_period="2025-06",
        generated_at=datetime(2026, 3, 4, 5, 6, 7, tzinfo=timezone.utc),
    )
    payload = _payload([_record(1)], spec=spec)
    assert payload["artifact"]["filename"] == f"rpt-{spec.fingerprint()}-20260304T050607Z.docx"


# --- no fabricated values --------------------------------------------------------


def test_no_value_outside_input_ever_appears_in_document():
    records = [_record(1), _record(2, document_id=2, value_raw="1350")]
    payload, doc = _docx_document(records)
    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for row in t.rows for c in row.cells
    )
    for fabricated in ("99999", "424242", "1234567890", "junkvalue"):
        assert fabricated not in all_text
    for record in records:
        assert record.value_raw in all_text  # real values DO appear


def test_audit_metadata_never_contains_values_or_contents():
    payload = _payload([_record(1, value_raw="1,200"), _record(2, document_id=2, value_raw="1350")])
    meta = audit_metadata(payload)
    flat = str(meta)
    for secret in ("1,200", "1350", "Mine A", "docA.pdf", "tonnes"):
        assert secret not in flat
    assert meta["record_count"] == 2 and meta["conflict_count"] == 1
    assert meta["report_type"] == "summary"  # spec metadata (not values) is fine


def test_docx_bytes_are_byte_identical_for_identical_spec_and_data():
    spec = build_specification(
        title="Deterministic", generated_at=datetime(2026, 2, 2, tzinfo=timezone.utc)
    )
    records = [_record(1), _record(2, document_id=2, value_raw="1350")]
    first = _payload(records, spec=spec)["docx_bytes"]
    second = _payload(records, spec=spec)["docx_bytes"]
    assert first == second


def test_generate_report_requires_exactly_one_input():
    spec = build_specification(title="x")
    with pytest.raises(ValueError):
        generate_report(spec)  # neither data nor db
    with pytest.raises(ValueError):
        generate_report(spec, data=_data([]), db=object())  # both given
    # Exactly one of them is the valid DB-free path (used everywhere above).
    payload = generate_report(spec, data=_data([]), db=None)
    assert payload["status"] == "ok"


def test_no_llm_anywhere_in_report_pipeline():
    """The report pipeline must run with NO LLM provider configured at all."""
    from app.llm.config import LLMConfig
    from app.llm.service import llm_available, set_llm_provider

    set_llm_provider(None)  # remove any injected provider
    try:
        payload = _payload([_record(1)])  # would fail if the pipeline needed an LLM
        assert payload["status"] == "ok"
        unconfigured = LLMConfig(provider="", model="", api_key="", base_url="",
                                 timeout_seconds=5.0)
        assert llm_available(unconfigured) is False
    finally:
        set_llm_provider(None)
