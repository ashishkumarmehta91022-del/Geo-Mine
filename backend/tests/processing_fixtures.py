"""Generated test fixtures for the processing pipeline (Step 4).

Every fixture is produced in-memory at test time with the same libraries the
extractors use — no binary blobs committed to the repo.
"""

import io
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from PIL import Image

# --- PDF ------------------------------------------------------------------


def pdf_bytes(pages: list[str]) -> bytes:
    """A text PDF with one page per given string (empty string = no text)."""
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_textbox(page.rect, text, fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data


def rendered_text_image_bytes(text: str, size: tuple[int, int] = (700, 180)) -> bytes:
    """A PNG containing rendered text — a deterministic stand-in for a scan."""
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.load_default(size=52)
    except TypeError:  # older Pillow without sized default font
        font = ImageFont.load_default()
    draw.text((24, size[1] // 3), text, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def pdf_bytes_with_image_pages(page_specs: list[str | None]) -> bytes:
    """Mixed PDF: each entry is native page text, or None for a scanned
    (image-only) page built from a rendered-text image."""
    import pymupdf

    doc = pymupdf.open()
    for spec in page_specs:
        page = doc.new_page()
        if spec is None:
            page.insert_image(page.rect, stream=rendered_text_image_bytes("SCANNED PAGE 12345"))
        else:
            page.insert_textbox(page.rect, spec, fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data


# --- DOCX --------------------------------------------------------------------


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        rows = len(table)
        cols = len(table[0]) if rows else 0
        doc_table = doc.add_table(rows=rows, cols=cols)
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                doc_table.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- XLSX ---------------------------------------------------------------------


def xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- XLS (legacy binary) ----------------------------------------------------------


def xls_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    import xlwt

    book = xlwt.Workbook()
    for sheet_name, rows in sheets.items():
        sheet = book.add_sheet(sheet_name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is None:
                    continue
                if isinstance(value, datetime):
                    sheet.write(r, c, value)
                elif isinstance(value, bool):
                    sheet.write(r, c, bool(value))
                elif isinstance(value, int | float):
                    sheet.write(r, c, value)
                else:
                    sheet.write(r, c, str(value))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- images ----------------------------------------------------------------------


def png_bytes(width: int = 64, height: int = 48) -> bytes:
    image = Image.new("RGB", (width, height), color=(30, 60, 120))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def jpeg_bytes(width: int = 64, height: int = 48) -> bytes:
    image = Image.new("RGB", (width, height), color=(120, 30, 60))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG")
    return buffer.getvalue()


# --- corrupted files ------------------------------------------------------------------


def corrupted_pdf_bytes() -> bytes:
    return b"%PDF-1.7 this file is truncated and has no xref table \x00\x01\x02"


def corrupted_docx_bytes() -> bytes:
    # Starts like a ZIP (passes naive magic checks) but is not a valid archive.
    return b"PK\x03\x04" + b"\x00" * 256


def corrupted_xlsx_bytes() -> bytes:
    return b"PK\x03\x04" + b"\xff" * 256
